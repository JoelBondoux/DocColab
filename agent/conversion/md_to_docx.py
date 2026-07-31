from __future__ import annotations

import io
from dataclasses import dataclass

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from markdown_it import MarkdownIt
from markdown_it.token import Token


@dataclass(slots=True)
class InlineStyle:
    bold: bool = False
    italic: bool = False
    code: bool = False
    href: str | None = None


def markdown_to_docx(markdown: str) -> bytes:
    """Pure-Python Markdown fallback for headings, lists, tables, emphasis, and links."""
    document = Document()
    _ensure_styles(document)
    tokens = MarkdownIt("commonmark").enable("table").parse(markdown)
    list_stack: list[str] = []
    blockquote_depth = 0
    index = 0
    while index < len(tokens):
        token = tokens[index]
        if token.type == "heading_open":
            level = int(token.tag[1])
            inline = tokens[index + 1]
            paragraph = document.add_heading(level=level)
            _render_inline(paragraph, inline.children or [])
            index += 3
            continue
        if token.type == "bullet_list_open":
            list_stack.append("List Bullet")
        elif token.type == "ordered_list_open":
            list_stack.append("List Number")
        elif token.type in {"bullet_list_close", "ordered_list_close"} and list_stack:
            list_stack.pop()
        elif token.type == "blockquote_open":
            blockquote_depth += 1
        elif token.type == "blockquote_close":
            blockquote_depth = max(0, blockquote_depth - 1)
        elif token.type == "paragraph_open":
            inline = tokens[index + 1]
            style = (
                list_stack[-1]
                if list_stack
                else ("Intense Quote" if blockquote_depth else None)
            )
            paragraph = document.add_paragraph(style=style)
            _render_inline(paragraph, inline.children or [])
            index += 3
            continue
        elif token.type in {"fence", "code_block"}:
            paragraph = document.add_paragraph(style="DocColab Code")
            paragraph.add_run(token.content.rstrip("\n"))
        elif token.type == "table_open":
            index = _render_table(document, tokens, index)
            continue
        elif token.type == "hr":
            paragraph = document.add_paragraph()
            paragraph.add_run("―" * 20)
        index += 1
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _render_inline(paragraph: Paragraph, tokens: list[Token]) -> None:
    style = InlineStyle()
    for token in tokens:
        if token.type == "strong_open":
            style.bold = True
        elif token.type == "strong_close":
            style.bold = False
        elif token.type == "em_open":
            style.italic = True
        elif token.type == "em_close":
            style.italic = False
        elif token.type == "link_open":
            attributes = dict(token.attrs or [])
            href = attributes.get("href")
            style.href = str(href) if href is not None else None
        elif token.type == "link_close":
            style.href = None
        elif token.type == "code_inline":
            _add_text(paragraph, token.content, style, code=True)
        elif token.type == "text":
            _add_text(paragraph, token.content, style)
        elif token.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()
        elif token.type == "image":
            alt = token.content or dict(token.attrs or []).get("alt", "image")
            _add_text(paragraph, f"[{alt}]", style)


def _add_text(
    paragraph: Paragraph,
    text: str,
    style: InlineStyle,
    *,
    code: bool = False,
) -> None:
    if style.href:
        _add_hyperlink(paragraph, text, style.href, style.bold, style.italic)
        return
    run = paragraph.add_run(text)
    run.bold = style.bold
    run.italic = style.italic
    if code or style.code:
        run.font.name = "Courier New"


def _add_hyperlink(
    paragraph: Paragraph,
    text: str,
    url: str,
    bold: bool,
    italic: bool,
) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    if bold:
        properties.append(OxmlElement("w:b"))
    if italic:
        properties.append(OxmlElement("w:i"))
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _render_table(
    document: DocumentObject,
    tokens: list[Token],
    start: int,
) -> int:
    rows: list[list[list[Token]]] = []
    current_row: list[list[Token]] | None = None
    index = start + 1
    while index < len(tokens) and tokens[index].type != "table_close":
        token = tokens[index]
        if token.type == "tr_open":
            current_row = []
        elif token.type == "tr_close" and current_row is not None:
            rows.append(current_row)
            current_row = None
        elif token.type in {"th_open", "td_open"} and current_row is not None:
            inline = tokens[index + 1] if index + 1 < len(tokens) else None
            current_row.append(
                (inline.children or []) if inline and inline.type == "inline" else []
            )
        index += 1
    if not rows:
        return index + 1
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, inline_tokens in enumerate(row):
            paragraph = table.cell(row_index, column_index).paragraphs[0]
            _render_inline(paragraph, inline_tokens)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    return index + 1


def _ensure_styles(document: DocumentObject) -> None:
    styles = document.styles
    if "DocColab Code" not in styles:
        style = styles.add_style("DocColab Code", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Courier New"
        style.font.size = Pt(9)
